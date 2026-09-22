"""
Envoi de mails personnalisés (cold outreach B2B) — Equation
------------------------------------------------------------
MULTI-EXPÉDITEUR : chaque commercial se connecte avec sa propre adresse +
mot de passe OVH. Le mail part réellement de sa boîte (From = lui, réponses
vers lui, alignement SPF/DKIM/DMARC du domaine). Aucun mot de passe n'est
stocké sur disque : il vit uniquement en mémoire le temps de la session.

Flask + SMTP OVH Email Pro. Construction MIME multipart (HTML + texte alterné
+ PJ + images inline cid:). Throttling, journal, accès par liste blanche.

IMPORTANT (Render) : UN SEUL worker gunicorn (-w 1 --threads 8). Les sessions
et les jobs vivent en mémoire ; plusieurs workers casseraient l'état.
"""

import os
import re
import time
import json
import uuid
import base64
import mimetypes
import smtplib
import imaplib
import tempfile
import threading
import urllib.request
import urllib.error
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email.mime.image import MIMEImage
from email.utils import formataddr
from email import encoders
from functools import wraps

from flask import (
    Flask, Request, request, session, redirect, url_for,
    render_template, jsonify, send_file, flash
)
from werkzeug.utils import secure_filename

import warmup  # moteur de cadencement & chauffe (warm-up)
import levees_auto  # brique « Levées +12 mois » (route manuelle ci-dessous ; le déclencheur hebdo vit dans warmup._loop)

# --- Sonde de diagnostic : piles d'appel de TOUS les threads sur SIGUSR1 ----
# Aucun effet en fonctionnement normal. En cas de blocage (threads en attente
# de verrou), « kill -USR1 <pid du worker> » écrit les piles Python complètes
# dans les logs Render, ce qui identifie précisément le verrou fautif.
try:
    import faulthandler as _fh
    import signal as _sig
    _fh.register(_sig.SIGUSR1, all_threads=True)
except Exception:
    pass

# --------------------------------------------------------------------------- #
# Configuration (variables d'environnement)
# --------------------------------------------------------------------------- #
SMTP_HOST = os.environ.get("SMTP_HOST", "pro1.mail.ovh.net")   # infra partagée
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587"))

# Copie dans le dossier "Envoyés" via IMAP (même serveur OVH que le SMTP).
SAVE_TO_SENT = os.environ.get("SAVE_TO_SENT", "1") not in ("0", "", "false", "False")
IMAP_HOST = os.environ.get("IMAP_HOST", SMTP_HOST)             # = SMTP_HOST par défaut
IMAP_PORT = int(os.environ.get("IMAP_PORT", "993"))
SENT_FOLDER = os.environ.get("SENT_FOLDER", "")               # vide = détection auto
REPLY_TO  = os.environ.get("REPLY_TO", "")                     # optionnel (global)

# Journal des envois dans Notion (optionnel) : renseigne NOTION_TOKEN +
# NOTION_DATABASE_ID dans Render pour activer. Sinon, ignoré silencieusement.
NOTION_TOKEN = os.environ.get("NOTION_TOKEN", "").strip()
NOTION_DATABASE_ID = os.environ.get("NOTION_DATABASE_ID", "6b7a962f02ef4207a0fdae7f244dd656").strip()
SECRET_KEY = os.environ.get("SECRET_KEY", os.urandom(24).hex())

# Liste blanche des expéditeurs autorisés (séparés par des virgules).
# Vide = toute adresse OVH valide peut se connecter. Recommandé : la remplir.
ALLOWED_SENDERS = {
    e.strip().lower()
    for e in os.environ.get("ALLOWED_SENDERS", "").split(",")
    if e.strip()
}

# Noms affichés (From). Modifiable librement. Fallback = partie avant @.
SENDER_NAMES = {
    "mbureau@equation-sie.com":      "Marine Bureau de Rotalier",
    "lbastian@equation-sie.com":     "Lionel Bastian",
    "rabou-khalil@equation-sie.com": "Richard Abou Khalil",
    "nvial@equation-sie.com":        "Nicolas Vial",
    "mmoriceau@equation-sie.com":    "Marie Moriceau",
    "smoinet@equation-sie.com":      "Séverine Moinet",
    "mbastian@equation-sie.com":     "Michel Bastian",
}

# Signatures HTML par expéditeur (fichier dédié signatures.py, facile à éditer).
try:
    from signatures import SIGNATURES
except ImportError:
    SIGNATURES = {}


MAX_TOTAL_ATTACH_MB = float(os.environ.get("MAX_TOTAL_ATTACH_MB", "25"))
COMPRESS_PDF_OVER_MB = float(os.environ.get("COMPRESS_PDF_OVER_MB", "4"))  # auto-compression au-delà
DEFAULT_DELAY_S     = int(os.environ.get("DEFAULT_DELAY_S", "45"))
OVH_HOURLY_LIMIT    = 200  # mails / heure / compte (doc OVH)

LOG_DIR  = os.environ.get("LOG_DIR", os.path.join(tempfile.gettempdir(), "relances_data"))
os.makedirs(LOG_DIR, exist_ok=True)
LOG_FILE = os.path.join(LOG_DIR, "journal_envois.jsonl")

# Mini-hébergeur de photos : images stockées sur le disque persistant, servies
# publiquement via /img/<nom> pour être chargées dans les mails (URL durable).
PHOTOS_DIR = os.path.join(LOG_DIR, "photos")
os.makedirs(PHOTOS_DIR, exist_ok=True)
ALLOWED_PHOTO = {".png", ".jpg", ".jpeg", ".gif", ".webp"}

ALLOWED_ATTACH = {".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp",
                  ".doc", ".docx", ".xls", ".xlsx", ".csv", ".txt"}
ALLOWED_INLINE = {".png", ".jpg", ".jpeg", ".gif", ".webp"}

app = Flask(__name__, template_folder=".")
app.secret_key = SECRET_KEY
# Marge au-dessus du plafond des PJ (en-têtes, encodage, plusieurs fichiers)
app.config["MAX_CONTENT_LENGTH"] = int((MAX_TOTAL_ATTACH_MB + 20) * 1024 * 1024)

# Au-delà de ce délai sans accès base réussi, /healthz répond 503 et Render
# redémarre l'instance (auto-guérison). Large pour éviter tout faux positif :
# le planificateur sonde la base toutes les ~20 s.
HEALTH_MAX_STALE_S = int(os.environ.get("HEALTH_MAX_STALE_S", "180"))


# --- Grandes listes de prospects collées dans le formulaire ---------------
# Werkzeug >= 3.1 plafonne les CHAMPS TEXTE d'un formulaire à 500 Ko par défaut.
# Or une liste Excel est collée en texte dans « Liste collée » : 678 contacts x
# 30 colonnes = ~540 Ko -> dépassement -> erreur 413 affichée à tort comme
# « Fichiers trop volumineux » alors qu'aucune pièce jointe n'est envoyée.
# On relève donc explicitement ce plafond (les PJ restent bornées par
# MAX_CONTENT_LENGTH ci-dessus et par la vérification métier des envois).
class _BigFormRequest(Request):
    max_form_memory_size = int((MAX_TOTAL_ATTACH_MB + 20) * 1024 * 1024)  # idem MAX_CONTENT_LENGTH
    max_form_parts = 10000                                                # marge sur le nb de champs


app.request_class = _BigFormRequest

# État en mémoire (d'où le worker unique)
# Identité (email, name) -> stockée dans le cookie signé (survit aux redémarrages).
# Mot de passe OVH -> uniquement en mémoire ici, jamais sur disque ni dans le cookie.
PASSWORDS = {}         # email -> mot de passe (perdu au redémarrage -> reconnexion)
JOBS = {}
LOCK = threading.Lock()

# --- Auto-login optionnel ---------------------------------------------------- #
# Pré-charge des mots de passe au démarrage pour éviter de se reconnecter après
# chaque redémarrage. À DÉFINIR UNIQUEMENT dans l'onglet Environment de Render,
# JAMAIS dans le code/GitHub.
#   - Cas simple (ton compte) : AUTOLOGIN_EMAIL + AUTOLOGIN_PASSWORD
#   - Plusieurs comptes        : AUTOLOGIN = "email1:mdp1;email2:mdp2"
# La connexion reste protégée par le cookie : un mot de passe pré-chargé ne
# connecte personne tout seul, il évite juste d'avoir à le re-saisir.
def _preload_passwords():
    # Compte simple : AUTOLOGIN_EMAIL + AUTOLOGIN_PASSWORD
    em = os.environ.get("AUTOLOGIN_EMAIL", "").strip().lower()
    pw = os.environ.get("AUTOLOGIN_PASSWORD", "")
    if em and pw:
        PASSWORDS[em] = pw
    # Plusieurs comptes (commerciaux) : paires numérotées, robustes aux caractères
    # spéciaux dans les mots de passe.
    #   AUTOLOGIN_1_EMAIL / AUTOLOGIN_1_PASSWORD, AUTOLOGIN_2_EMAIL / ...
    for n in range(1, 21):
        e = os.environ.get("AUTOLOGIN_%d_EMAIL" % n, "").strip().lower()
        p = os.environ.get("AUTOLOGIN_%d_PASSWORD" % n, "")
        if e and p:
            PASSWORDS[e] = p
    # Format compact optionnel : AUTOLOGIN = "email1:mdp1;email2:mdp2"
    blob = os.environ.get("AUTOLOGIN", "").strip()
    for pair in blob.split(";"):
        if ":" in pair:
            e, p = pair.split(":", 1)
            e = e.strip().lower(); p = p.strip()
            if e and p:
                PASSWORDS[e] = p
_preload_passwords()


def _password_for(email):
    """Mot de passe en mémoire pour cette boîte. S'IL MANQUE (ex. après une
    déconnexion / un changement de compte, qui purge la mémoire via PASSWORDS.pop,
    ou un process qui n'aurait pas préchargé), on RECHARGE d'abord depuis AUTOLOGIN.
    Ainsi une boîte configurée dans AUTOLOGIN ne peut jamais rester durablement sans
    mot de passe — c'était la cause de l'inondation « pas de mot de passe » du 24/07."""
    p = PASSWORDS.get(email)
    if p is None:
        _preload_passwords()          # relit AUTOLOGIN (idempotent, aucune valeur effacée)
        p = PASSWORDS.get(email)
    return p


# --------------------------------------------------------------------------- #
# Authentification (identifiants OVH de chacun)
# --------------------------------------------------------------------------- #
def display_name(email: str) -> str:
    return SENDER_NAMES.get(email.lower(), email.split("@")[0].replace(".", " ").title())


def smtp_check(email: str, password: str):
    """Valide les identifiants en se connectant réellement à OVH."""
    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as s:
            s.ehlo(); s.starttls(); s.ehlo()
            s.login(email, password)
        return True, None
    except smtplib.SMTPAuthenticationError:
        return False, "Adresse ou mot de passe OVH incorrect."
    except Exception as e:
        return False, f"Connexion au serveur impossible : {e}"


def current_user():
    """Identité issue du cookie signé (survit aux redémarrages)."""
    email = session.get("email")
    if not email:
        return None
    return {"email": email, "name": session.get("name", display_name(email))}


def has_password(email):
    return bool(PASSWORDS.get(email))


@app.context_processor
def inject_user():
    return {"current_user": current_user()}


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        u = current_user()
        if not u:
            return redirect(url_for("login"))
        # Mot de passe perdu (redémarrage du service) -> reconnexion immédiate,
        # plutôt que de laisser composer puis bloquer à l'envoi.
        if not has_password(u["email"]):
            flash("Session expirée (le service a redémarré). Reconnecte-toi pour continuer.")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapper


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        if ALLOWED_SENDERS and email not in ALLOWED_SENDERS:
            flash("Cette adresse n'est pas autorisée à utiliser l'outil.")
            return render_template("login.html")
        ok, err = smtp_check(email, password)
        if not ok:
            flash(err)
            return render_template("login.html")
        session["email"] = email
        session["name"] = display_name(email)
        PASSWORDS[email] = password
        return redirect(url_for("accueil"))
    return render_template("login.html")


@app.route("/logout")
def logout():
    email = session.get("email")
    PASSWORDS.pop(email, None)
    session.clear()
    return redirect(url_for("login"))


# --------------------------------------------------------------------------- #
# Outils : personnalisation, parsing, conversion texte
# --------------------------------------------------------------------------- #
# Variables {colonne} : accepte lettres accentuées, chiffres, _ mais aussi
# ESPACES, tirets et apostrophes à l'intérieur (ex. {adresse prospect}).
# La classe interne exclut { } donc on ne franchit jamais une accolade.
FIELD_RE = re.compile(r"\{(\w[\w \-']*\w|\w)\}")


def personalize(template: str, row: dict) -> str:
    def sub(m):
        key = m.group(1)
        if key in row:
            return str(row.get(key, ""))
        # tolérance : la clé pourrait être stockée sans espaces superflus / en minuscules
        k2 = key.strip().lower()
        if k2 in row:
            return str(row.get(k2, ""))
        return m.group(0)          # variable inconnue -> laissée telle quelle
    return FIELD_RE.sub(sub, template or "")


def parse_recipients(raw: str):
    raw = (raw or "").strip()
    if not raw:
        return [], ["Liste de destinataires vide."]
    lines = [l for l in raw.splitlines() if l.strip()]

    # MODE SIMPLE : si la 1re ligne contient déjà un "@", c'est une liste
    # d'adresses (une par ligne et/ou séparées par des virgules), pas un en-tête.
    if "@" in lines[0]:
        rows, errors, seen = [], [], set()
        for tok in re.split(r"[\s,;]+", raw):
            tok = tok.strip()
            if not tok:
                continue
            if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", tok):
                errors.append(f"Adresse ignorée (invalide) : {tok}")
                continue
            if tok.lower() in seen:
                continue
            seen.add(tok.lower())
            rows.append({"email": tok})
        if not rows:
            errors.append("Aucune adresse e-mail valide trouvée.")
        return rows, errors

    # MODE TABLEAU : en-tête avec colonnes (pour personnalisation {prenom}, etc.)
    sep = "\t" if "\t" in lines[0] else (";" if ";" in lines[0] else ",")
    headers = [h.strip().lower() for h in lines[0].split(sep)]
    if "email" not in headers:
        return [], ["Colle des adresses (une par ligne ou séparées par des virgules), "
                    "ou un tableau dont la 1re ligne contient une colonne 'email'."]
    rows, errors, seen = [], [], set()
    for i, line in enumerate(lines[1:], start=2):
        cells = [c.strip() for c in line.split(sep)]
        row = {headers[j]: (cells[j] if j < len(cells) else "")
               for j in range(len(headers))}
        email = row.get("email", "").strip()
        if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
            errors.append(f"Ligne {i} : email invalide ({email or 'vide'}).")
            continue
        if email.lower() in seen:
            errors.append(f"Ligne {i} : doublon ignoré ({email}).")
            continue
        seen.add(email.lower())
        rows.append(row)
    return rows, errors


def html_to_text(html: str) -> str:
    txt = html or ""
    txt = re.sub(r"(?is)<br\s*/?>", "\n", txt)
    txt = re.sub(r"(?is)</p>", "\n\n", txt)
    txt = re.sub(r'(?is)<a[^>]*href="([^"]+)"[^>]*>(.*?)</a>',
                 lambda m: f"{re.sub('<[^>]+>', '', m.group(2))} ({m.group(1)})", txt)
    txt = re.sub(r"(?is)<[^>]+>", "", txt)
    txt = (txt.replace("&nbsp;", " ").replace("&amp;", "&")
              .replace("&lt;", "<").replace("&gt;", ">").replace("&#39;", "'"))
    return re.sub(r"\n{3,}", "\n\n", txt).strip()


# --------------------------------------------------------------------------- #
# Photos dans le corps : insertion auto + rendu dans l'aperçu
# --------------------------------------------------------------------------- #
PHOTO_TAG = ('<img src="cid:{cid}" alt="{name}" '
             'style="max-width:100%;height:auto;display:block;margin:14px 0;border-radius:4px;">')


def photo_tag(cid, name, center=False):
    """Toutes les photos à la même largeur (600px, réduites si l'écran est plus
    petit). Centrées ou alignées à gauche selon le choix."""
    margin = "14px auto" if center else "14px 0"
    return ('<img src="cid:%s" alt="%s" '
            'style="width:600px;max-width:100%%;height:auto;display:block;'
            'margin:%s;border-radius:4px;">' % (cid, name, margin))


def embed_photos(body: str, inline_images, center=False) -> str:
    """Place chaque photo là où un repère apparaît dans le corps. Repères tolérés
    (espaces et majuscules ignorés) : [photo1], [photo 1], [Photo1], ou [nomdufichier].
    Sinon (aucun repère, ni cid: manuel), la photo est ajoutée à la fin."""
    body = body or ""
    for i, img in enumerate(inline_images, start=1):
        tag = photo_tag(img["cid"], img["filename"], center)
        base = os.path.splitext(img["filename"])[0]
        patterns = [
            r"\[\s*photo\s*" + str(i) + r"\s*\]",                 # [photo1] / [ photo 1 ]
            r"\[\s*" + re.escape(base) + r"\s*\]",                # [nomdufichier]
            r"\[\s*" + re.escape(img["filename"]) + r"\s*\]",     # [nomdufichier.png]
        ]
        for pat in patterns:
            body = re.sub(pat, lambda _m: tag, body, flags=re.IGNORECASE)
    # Filet de sécurité : photos non placées -> à la fin
    extra = "".join(
        photo_tag(img["cid"], img["filename"], center)
        for img in inline_images
        if f"cid:{img['cid']}" not in body
    )
    return body + extra


def nl2br(text: str) -> str:
    """Convertit les sauts de ligne tapés (touche Entrée) en vrais retours HTML,
    pour que le texte s'affiche comme à l'écran. Ne casse pas les liens (inline)
    ni le HTML déjà présent : on évite les doubles sauts autour des blocs."""
    if not text:
        return text
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = t.replace("\n", "<br>\n")
    # pas de <br> superflu juste après/avant un bloc HTML
    t = re.sub(r"(</(?:p|div|ul|ol|li|h[1-6]|blockquote|table|tr)>)\s*<br>", r"\1", t, flags=re.I)
    t = re.sub(r"<br>\s*(<(?:p|div|ul|ol|li|h[1-6]|blockquote|table|tr)[ >])", r"\1", t, flags=re.I)
    return t


# Image DIFFÉRENTE par prospect : marqueur [img:nom_de_colonne] -> <img src="URL de la ligne">
IMG_MARK = re.compile(r"\[img:\s*([^\]\n]+?)\s*\]", re.I)


def insert_row_images(body: str, row: dict, center=False, resolver=None) -> str:
    """Remplace [img:colonne] par l'image propre à chaque destinataire.
    `resolver(col, url)` décide la source finale de l'image :
      - à l'envoi : renvoie 'cid:...' (image embarquée dans le mail) ;
      - à l'aperçu : renvoie une data-URI (image téléchargée à la volée) ;
      - par défaut (None) : le lien direct s'il est http(s), sinon rien.
    Si la source est vide/invalide, le marqueur est simplement retiré."""
    def repl(m):
        col = (m.group(1) or "").strip().lower()
        url = str(row.get(col, "") or "").strip()
        if resolver is not None:
            src = resolver(col, url) or ""
        else:
            src = url if re.match(r"^https?://|^data:image", url, re.I) else ""
        if not src:
            return ""
        safe = src.replace('"', "%22")
        if src.startswith("http"):
            safe = safe.replace("<", "").replace(">", "").replace(" ", "%20")
        margin = "14px auto" if center else "14px 0"
        return ('<img src="%s" alt="" style="width:600px;max-width:100%%;height:auto;'
                'display:block;margin:%s;border-radius:4px;">' % (safe, margin))
    return IMG_MARK.sub(repl, body or "")


def assemble_body(template: str, row: dict, photos, signature_html: str, center=False,
                  img_resolver=None) -> str:
    """Corps final : texte personnalisé -> images par prospect -> photos -> signature."""
    body = personalize(template, row)
    body = insert_row_images(body, row, center, resolver=img_resolver)   # image propre à la ligne
    body = nl2br(body)                            # touche Entrée = vrai retour à la ligne
    body = embed_photos(body, photos, center)
    if signature_html:
        body = body + "<br>" + signature_html
    return body


def inline_to_data(body: str, photos) -> str:
    """Rend les photos cid: visibles dans l'aperçu (base64). La signature
    utilise une image en ligne (URL) qui s'affiche déjà toute seule."""
    for img in photos:
        try:
            with open(img["path"], "rb") as fh:
                b64 = base64.b64encode(fh.read()).decode()
            mime = mimetypes.guess_type(img["filename"])[0] or "image/png"
            body = body.replace(f"cid:{img['cid']}", f"data:{mime};base64,{b64}")
        except OSError:
            pass
    return body


# --------------------------------------------------------------------------- #
# Compression PDF (ré-échantillonne les images). Best-effort, optionnelle.
# --------------------------------------------------------------------------- #
def compress_pdf(in_path, max_dim=1600, quality=72):
    """Réduit le poids d'un PDF trop lourd en ré-échantillonnant ses images.
    Retourne (chemin, nouvelle_taille) ; si échec/inutile -> (in_path, None)."""
    try:
        import fitz
        from PIL import Image
        import io as _io
    except ImportError:
        return in_path, None
    try:
        doc = fitz.open(in_path)
        for page in doc:
            for info in page.get_images(full=True):
                xref = info[0]
                try:
                    pix = fitz.Pixmap(doc, xref)
                    if pix.n >= 5:                       # CMYK / alpha -> RGB
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    w, h = pix.width, pix.height
                    mode = "RGB" if pix.n >= 3 else "L"
                    img = Image.frombytes(mode, (w, h), pix.samples)
                    scale = min(1.0, max_dim / max(w, h))
                    if scale < 1.0:
                        img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))),
                                         Image.LANCZOS)
                    buf = _io.BytesIO()
                    img.convert("RGB").save(buf, "JPEG", quality=quality)
                    doc.update_stream(xref, buf.getvalue())
                    doc.xref_set_key(xref, "Filter", "/DCTDecode")
                    doc.xref_set_key(xref, "Width", str(img.width))
                    doc.xref_set_key(xref, "Height", str(img.height))
                    doc.xref_set_key(xref, "BitsPerComponent", "8")
                    doc.xref_set_key(xref, "ColorSpace", "/DeviceRGB")
                    pix = None
                except Exception:
                    continue
        out = in_path + ".min.pdf"
        doc.save(out, garbage=4, deflate=True, clean=True)
        doc.close()
        if os.path.exists(out) and os.path.getsize(out) < os.path.getsize(in_path):
            return out, os.path.getsize(out)
        if os.path.exists(out):
            os.remove(out)
        return in_path, None
    except Exception:
        return in_path, None


# --------------------------------------------------------------------------- #
# Construction du message MIME
# --------------------------------------------------------------------------- #
def build_message(from_email, from_name, to_email, subject,
                  html_body, attachments, inline_images):
    root = MIMEMultipart("mixed")
    root["Subject"] = subject
    root["From"] = formataddr((from_name, from_email))
    root["To"] = to_email
    if REPLY_TO:
        root["Reply-To"] = REPLY_TO

    related = MIMEMultipart("related")
    alt = MIMEMultipart("alternative")
    alt.attach(MIMEText(html_to_text(html_body), "plain", "utf-8"))
    alt.attach(MIMEText(html_body, "html", "utf-8"))
    related.attach(alt)

    for img in inline_images:
        with open(img["path"], "rb") as fh:
            raw = fh.read()
        try:
            part = MIMEImage(raw)                       # devine le format (imghdr)
        except Exception:
            # Repli si le format n'est pas devinable : on prend l'extension du fichier.
            sub = os.path.splitext(img.get("filename", ""))[1].lstrip(".").lower() or "jpeg"
            if sub == "jpg":
                sub = "jpeg"
            part = MIMEImage(raw, _subtype=sub)
        part.add_header("Content-ID", f"<{img['cid']}>")
        part.add_header("Content-Disposition", "inline", filename=img["filename"])
        related.attach(part)
    root.attach(related)

    for att in attachments:
        with open(att["path"], "rb") as fh:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(fh.read())
        encoders.encode_base64(part)
        part.add_header("Content-Disposition", "attachment", filename=att["filename"])
        root.attach(part)
    return root


# --------------------------------------------------------------------------- #
# Copie dans le dossier "Envoyés" (IMAP)
# --------------------------------------------------------------------------- #
def _sent_candidates(M):
    """Liste ordonnée de dossiers 'Envoyés' à essayer : flag \\Sent d'abord,
    puis noms contenant sent/envoy, puis valeurs par défaut courantes OVH."""
    cands = []
    if SENT_FOLDER:
        cands.append(SENT_FOLDER)
    try:
        typ, data = M.list()
        if typ == "OK" and data:
            flagged, named = [], []
            for raw in data:
                line = raw.decode("ascii", "ignore") if isinstance(raw, bytes) else str(raw)
                quoted = re.findall(r'"([^"]*)"', line)
                name = quoted[-1] if quoted else line.split()[-1].strip('"')
                if "\\Sent" in line:
                    flagged.append(name)
                elif re.search(r"(?i)(sent|envoy)", name):
                    named.append(name)
            cands += flagged + named
    except Exception:
        pass
    # Valeurs par défaut courantes (OVH / Open-Xchange)
    cands += ["Sent", "Envoy&AOk-s", "INBOX.Sent", "INBOX.Envoy&AOk-s",
              "Sent Messages", "INBOX.Sent Messages"]
    seen, ordered = set(), []
    for c in cands:
        if c and c not in seen:
            seen.add(c); ordered.append(c)
    return ordered


def save_to_sent(email, password, msg):
    """Range une copie du message dans le dossier Envoyés. Best-effort.
    Essaie plusieurs noms de dossier ; renvoie la dernière erreur si tout échoue."""
    if not SAVE_TO_SENT:
        return True, None
    last_err = "aucun dossier Envoyés trouvé"
    try:
        M = imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT, timeout=20)
        try:
            M.login(email, password)
            when = imaplib.Time2Internaldate(time.time())
            raw = msg.as_bytes()
            for folder in _sent_candidates(M):
                try:
                    typ, _ = M.append('"%s"' % folder, "\\Seen", when, raw)
                    if typ == "OK":
                        return True, None
                    last_err = "réponse %s pour le dossier %s" % (typ, folder)
                except Exception as e:
                    last_err = "%s (dossier %s)" % (e, folder)
            # Aucun dossier "Envoyés" trouvé — fréquent sur une boîte NEUVE (aucun mail
            # n'y a encore été classé). On tente d'en créer un et d'y déposer la copie.
            for folder in ("Sent", "Envoy&AOk-s"):
                try:
                    M.create('"%s"' % folder)
                except Exception:
                    pass
                try:
                    try:
                        M.subscribe('"%s"' % folder)
                    except Exception:
                        pass
                    typ, _ = M.append('"%s"' % folder, "\\Seen", when, raw)
                    if typ == "OK":
                        return True, None
                    last_err = "réponse %s (dossier créé %s)" % (typ, folder)
                except Exception as e:
                    last_err = "création+dépôt %s : %s" % (folder, e)
        finally:
            try: M.logout()
            except Exception: pass
    except Exception as e:
        return False, "IMAP %s:%s — %s" % (IMAP_HOST, IMAP_PORT, e)
    return False, last_err
# --------------------------------------------------------------------------- #
def log_to_notion(commercial, destinataire, statut, objet="", copie_ok=False, detail=""):
    """Crée une ligne dans la base Notion (best-effort, ne bloque jamais l'envoi)."""
    if not (NOTION_TOKEN and NOTION_DATABASE_ID):
        return
    props = {
        "Destinataire": {"title": [{"text": {"content": (destinataire or "—")[:1900]}}]},
        "Statut": {"select": {"name": statut}},
        "Date d'envoi": {"date": {"start": datetime.now().isoformat(timespec="seconds")}},
        "Copie Envoyés": {"checkbox": bool(copie_ok)},
    }
    if objet:
        props["Objet"] = {"rich_text": [{"text": {"content": objet[:1900]}}]}
    if detail:
        props["Détail erreur"] = {"rich_text": [{"text": {"content": detail[:1900]}}]}
    if commercial:
        props["Commercial"] = {"select": {"name": commercial[:100]}}
    payload = json.dumps({"parent": {"database_id": NOTION_DATABASE_ID},
                          "properties": props}).encode("utf-8")
    req = urllib.request.Request(
        "https://api.notion.com/v1/pages", data=payload, method="POST",
        headers={"Authorization": "Bearer " + NOTION_TOKEN,
                 "Notion-Version": "2022-06-28",
                 "Content-Type": "application/json"})
    try:
        urllib.request.urlopen(req, timeout=12)
    except urllib.error.HTTPError as e:
        try: detail_msg = e.read().decode("utf-8", "ignore")[:300]
        except Exception: detail_msg = ""
        print("[notion] échec HTTP %s : %s" % (e.code, detail_msg), flush=True)
    except Exception as e:
        print("[notion] échec : %s" % e, flush=True)


def log_send(entry: dict):
    entry["ts"] = datetime.now().isoformat(timespec="seconds")
    with open(LOG_FILE, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(entry, ensure_ascii=False) + "\n")


# --------------------------------------------------------------------------- #
# Thread d'envoi (throttling) — utilise les identifiants de l'expéditeur du job
# --------------------------------------------------------------------------- #
def run_job(job_id):
    with LOCK:
        job = JOBS[job_id]

    # Compression des gros PDF ici (en tâche de fond, ne bloque pas l'aperçu)
    if any(a.get("to_compress") for a in job["attachments"]):
        job["status"] = "préparation"
        for a in job["attachments"]:
            if a.get("to_compress"):
                try:
                    orig = os.path.getsize(a["path"])
                    newpath, newsize = compress_pdf(a["path"])
                    if newsize and newsize < orig:
                        a["path"] = newpath
                        a["size_kb"] = round(newsize / 1024)
                        a["orig_kb"] = round(orig / 1024)
                except Exception:
                    pass
                a["to_compress"] = False

    job["status"] = "running"

    s_email = job["sender_email"]
    s_pass  = job["sender_password"]
    s_name  = job["sender_name"]
    delay   = job["delay"]
    recipients = job["recipients"]

    # --- Copie "Envoyés" : UNE seule connexion IMAP réutilisée pour toute la
    # série (fiable même pour plusieurs mails rapprochés ; reconnexion si coupée).
    sent = {"M": None, "folders": None}
    def _imap_ok():
        if not SAVE_TO_SENT:
            return False
        if sent["M"] is not None:
            try:
                sent["M"].noop(); return True
            except Exception:
                sent["M"] = None
        try:
            M = imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT, timeout=20)
            M.login(s_email, s_pass)
            sent["M"] = M
            if sent["folders"] is None:
                sent["folders"] = _sent_candidates(M) or ["Sent"]
            return True
        except Exception:
            sent["M"] = None
            return False
    def _copy_sent(msg):
        if not SAVE_TO_SENT:
            return True, None
        last = "connexion IMAP impossible"
        for _ in range(2):                      # 1 essai + 1 reconnexion
            if not _imap_ok():
                continue
            when = imaplib.Time2Internaldate(time.time())
            raw = msg.as_bytes()
            for folder in sent["folders"]:
                try:
                    typ, _r = sent["M"].append('"%s"' % folder, "\\Seen", when, raw)
                    if typ == "OK":
                        sent["folders"] = [folder]   # fige le dossier qui marche
                        return True, None
                    last = "réponse %s (%s)" % (typ, folder)
                except Exception as e:
                    last = "%s (%s)" % (e, folder)
                    sent["M"] = None                 # force la reconnexion
                    break
        return False, last

    for idx, row in enumerate(recipients):
        if job.get("cancel"):
            job["status"] = "annulé"; break
        to_email = row["email"]
        subject = personalize(job["subject"], row)
        body = assemble_body(job["body"], row, job["inline_images"], job["signature_html"], job.get("center", False))
        result = {"email": to_email, "subject": subject, "expediteur": s_email}
        try:
            msg = build_message(s_email, s_name, to_email, subject, body,
                                job["attachments"], job["inline_images"])
            with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as server:
                server.ehlo(); server.starttls(); server.ehlo()
                server.login(s_email, s_pass)
                server.send_message(msg)
            result["status"] = "OK"
            # Copie dans Envoyés (n'échoue jamais l'envoi lui-même)
            copied, cerr = _copy_sent(msg)
            result["copie"] = "OK" if copied else "KO"
            if not copied:
                result["copie_err"] = cerr
        except Exception as e:
            result["status"] = "ERREUR"; result["error"] = str(e)

        log_send(dict(result))
        # Journal Notion (best-effort)
        detail = result.get("error") or result.get("copie_err") or ""
        log_to_notion(commercial=s_name, destinataire=to_email,
                      statut=("Envoyé" if result["status"] == "OK" else "Erreur"),
                      objet=subject, copie_ok=(result.get("copie") == "OK"), detail=detail)
        with LOCK:
            job["results"].append(result); job["done"] = idx + 1
            job["ok" if result["status"] == "OK" else "ko"] += 1

        if idx < len(recipients) - 1 and not job.get("cancel"):
            slept = 0
            while slept < delay and not job.get("cancel"):
                time.sleep(1); slept += 1

    if job["status"] != "annulé":
        job["status"] = "terminé"
    if sent["M"] is not None:               # ferme la connexion Envoyés
        try: sent["M"].logout()
        except Exception: pass
    job["sender_password"] = None  # on ne garde pas le mot de passe après usage
    tmp = job.get("tmpdir")
    if tmp and os.path.isdir(tmp):
        for f in os.listdir(tmp):
            try: os.remove(os.path.join(tmp, f))
            except OSError: pass


# --------------------------------------------------------------------------- #
# Routes
# --------------------------------------------------------------------------- #
@app.route("/")
@login_required
def accueil():
    """Écran d'accueil : choix du mode (envoi simple / campagne)."""
    u = current_user()
    return render_template("accueil.html", from_name=u["name"], from_email=u["email"])


@app.route("/composer")
@login_required
def index():
    u = current_user()
    # mode : 'simple' (1 mail) ou 'campagne' (séquence + relances)
    mode = request.args.get("mode", "campagne")
    if mode not in ("simple", "campagne"):
        mode = "campagne"
    pf = {}
    prefill_steps = []

    edit_id = request.args.get("edit", "")
    if edit_id:
        with LOCK:
            job = JOBS.get(edit_id)
        if job and job.get("sender_email") == u["email"]:
            pf = {
                "recipients": job.get("recipients_raw", ""),
                "name": job.get("name", ""),
                "subject": job.get("subject", ""),
                "body": job.get("body", ""),
                "delay": job.get("delay", DEFAULT_DELAY_S),
                "signature": job.get("want_sig", True),
                "center": job.get("center", False),
                "had_files": bool(job.get("attachments") or job.get("inline_images")),
                "start_at": (job.get("start_at") or "")[:16].replace(" ", "T"),
            }
            mode = job.get("mode", mode)
            # Restaure les RELANCES déjà écrites (pour ne rien reperdre via « ← Modifier »)
            jsteps = job.get("steps") or []
            if len(jsteps) > 1:
                prefill_steps = [{"delay_days": s.get("delay_days", 4),
                                  "subject": s.get("subject", ""),
                                  "body": s.get("body", "")} for s in jsteps[1:]]

    # Préremplissage depuis un modèle (mode campagne)
    tid = request.args.get("template", "")
    if tid:
        tpl = warmup.load_template_steps(tid)
        if tpl and tpl["steps"]:
            mode = "campagne"
            s0 = tpl["steps"][0]
            pf.setdefault("subject", s0.get("subject", ""))
            pf.setdefault("body", s0.get("body", ""))
            pf.setdefault("name", tpl.get("name", ""))
            prefill_steps = tpl["steps"][1:]   # relances à reconstruire

    return render_template("index.html", default_delay=DEFAULT_DELAY_S,
                           max_mb=MAX_TOTAL_ATTACH_MB, ovh_limit=OVH_HOURLY_LIMIT,
                           from_name=u["name"], from_email=u["email"], pf=pf,
                           mode=mode, prefill_steps=prefill_steps,
                           past_campaigns=warmup.all_campaigns())


@app.route("/aide")
@login_required
def aide():
    return render_template("aide.html")


@app.route("/guide")
def guide():
    """Guide « Lancer une campagne avec relances » — page PUBLIQUE (pas de connexion
    requise), pour la partager librement avec les commerciaux. Servie en brut pour
    éviter tout traitement Jinja (le guide contient des {prenom}, {adresse}…)."""
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "guide.html")
    try:
        with open(path, encoding="utf-8") as fh:
            return fh.read()
    except OSError:
        return "Guide indisponible pour le moment.", 404


@app.route("/img/<path:name>")
def serve_photo(name):
    """Sert une photo hébergée — PUBLIC (pas de connexion) pour que les clients mail
    puissent la charger. Anti-traversal : on ne sort jamais du dossier photos."""
    safe = os.path.basename(name)
    fp = os.path.join(PHOTOS_DIR, safe)
    if os.path.splitext(safe)[1].lower() not in ALLOWED_PHOTO or not os.path.isfile(fp):
        return "Image introuvable.", 404
    return send_file(fp, max_age=86400)


@app.route("/photos")
@login_required
def photos():
    imgs = []
    try:
        for fn in sorted(os.listdir(PHOTOS_DIR)):
            if os.path.splitext(fn)[1].lower() in ALLOWED_PHOTO:
                imgs.append(fn)
    except OSError:
        pass
    base = request.url_root.rstrip("/")   # ex. https://relances-chauffe.onrender.com
    return render_template("photos.html", images=imgs, base=base)


@app.route("/photos/upload", methods=["POST"])
@login_required
def photos_upload():
    files = request.files.getlist("images")
    saved = 0
    for f in files:
        if not f or not f.filename:
            continue
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ALLOWED_PHOTO:
            continue
        stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", os.path.splitext(f.filename)[0]).strip("-")[:40] or "photo"
        name = "%s-%s%s" % (stem, uuid.uuid4().hex[:8], ext)
        try:
            f.save(os.path.join(PHOTOS_DIR, name))
            saved += 1
        except Exception:
            pass
    if saved:
        flash("✅ %d photo(s) ajoutée(s). Copie l'URL de chacune dans la colonne Photo_url de ton Excel." % saved)
    else:
        flash("Aucune image valide (formats acceptés : jpg, png, gif, webp).")
    return redirect(url_for("photos"))


@app.route("/photos/<name>/delete", methods=["POST"])
@login_required
def photos_delete(name):
    safe = os.path.basename(name)
    fp = os.path.join(PHOTOS_DIR, safe)
    if os.path.isfile(fp) and os.path.splitext(safe)[1].lower() in ALLOWED_PHOTO:
        try:
            os.remove(fp)
            flash("🗑️ Photo supprimée.")
        except Exception:
            flash("Suppression impossible pour l'instant.")
    return redirect(url_for("photos"))


@app.route("/modele")
@login_required
def modele():
    """Génère et renvoie le modèle Excel des destinataires."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from io import BytesIO
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Destinataires"
    headers = ["email", "prenom", "societe", "adresse", "ville"]
    ws.append(headers)
    for r in [["jean.dupont@entreprise.fr", "Jean", "Entreprise SA", "12 rue de la Paix", "Paris 2e"],
              ["s.martin@societe-exemple.com", "Sophie", "Société Martin", "8 avenue Foch", "Lyon 6e"],
              ["contact@globex.fr", "Camille", "Globex", "3 place Bellecour", "Lyon 2e"]]:
        ws.append(r)
    n = len(headers)
    thin = Side(style="thin", color="DDDDDD"); border = Border(thin, thin, thin, thin)
    hf = Font(name="Arial", bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="3C214B")
    for c in range(1, n + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hf; cell.fill = hfill
        cell.alignment = Alignment(horizontal="left", vertical="center"); cell.border = border
    ws.row_dimensions[1].height = 22
    for rr in range(2, 5):
        for c in range(1, n + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name="Arial", italic=True, color="888888"); cell.border = border
    for i, w in enumerate([34, 16, 24, 30, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Lisez-moi")
    notes = ["Modèle — Liste de destinataires", "",
             "1. Remplace les lignes d'exemple (en gris) par tes vrais destinataires.",
             "2. Garde la 1re ligne d'en-têtes (email, prenom, societe, adresse, ville).",
             "3. La colonne « email » est obligatoire ; les autres sont optionnelles.",
             "4. Tu peux AJOUTER tes propres colonnes (ex. surface, loyer, reference) :",
             "   chaque colonne devient une variable utilisable dans le mail.",
             "5. Dans l'outil, glisse ce fichier dans la zone Destinataires (ou clique pour le choisir).",
             "",
             "Personnalisation : écris {prenom}, {adresse}, {ville}… dans ton mail,",
             "et chacun recevra sa version. Garde des en-têtes en minuscules, sans accent."]
    for i, t in enumerate(notes, start=1):
        cell = ws2.cell(row=i, column=1, value=t)
        cell.font = (Font(name="Arial", bold=True, size=14, color="3C214B") if i == 1
                     else Font(name="Arial"))
    ws2.column_dimensions["A"].width = 90
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="Modele-destinataires.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _links_new_tab(html):
    """Dans l'aperçu, tout lien s'ouvre dans un nouvel onglet."""
    def add(m):
        tag = m.group(0)
        if re.search(r"\btarget\s*=", tag, re.I):
            return tag
        return tag[:-1] + ' target="_blank" rel="noopener">'
    return re.sub(r"<a\b[^>]*>", add, html, flags=re.I)


def _parse_step(suffix, tmpdir, errors):
    """Lit un message (objet/corps/photos/PJ) depuis le formulaire, pour l'étape
    dont les champs portent le suffixe donné ("" pour l'étape 0, "_1" pour la
    relance 1, etc.). Renvoie un dict d'étape ou None si l'étape est vide."""
    subject = request.form.get("subject" + suffix, "").strip()
    body = request.form.get("body" + suffix, "").strip()
    if not subject and not body:
        return None  # étape non renseignée -> ignorée

    label = "Étape 1" if suffix == "" else ("Relance " + suffix.lstrip("_"))
    if not subject:
        errors.append("%s : objet manquant." % label)
    if not body:
        errors.append("%s : corps du mail manquant." % label)

    attachments, inline_images, total_bytes = [], [], 0
    for f in request.files.getlist("attachments" + suffix):
        if not f or not f.filename:
            continue
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ALLOWED_ATTACH:
            errors.append("%s : pièce jointe refusée (%s) : %s" % (label, ext, f.filename)); continue
        fname = secure_filename(f.filename)
        path = os.path.join(tmpdir, "att%s_%s" % (suffix, fname)); f.save(path)
        sz = os.path.getsize(path)
        entry = {"path": path, "filename": fname}
        if ext == ".pdf" and sz > COMPRESS_PDF_OVER_MB * 1024 * 1024:
            entry["to_compress"] = True
        total_bytes += sz; entry["size_kb"] = round(sz / 1024)
        attachments.append(entry)

    for f in request.files.getlist("inline" + suffix):
        if not f or not f.filename:
            continue
        ext = os.path.splitext(f.filename)[1].lower()
        if ext not in ALLOWED_INLINE:
            errors.append("%s : image refusée (%s) : %s" % (label, ext, f.filename)); continue
        fname = secure_filename(f.filename)
        path = os.path.join(tmpdir, "inl%s_%s" % (suffix, fname)); f.save(path)
        sz = os.path.getsize(path); total_bytes += sz
        inline_images.append({"path": path, "filename": fname, "cid": fname,
                              "size_kb": round(sz / 1024)})

    total_mb = total_bytes / (1024 * 1024)
    if total_mb > MAX_TOTAL_ATTACH_MB:
        errors.append("%s : pièces jointes trop lourdes (%.1f Mo > %s Mo)."
                      % (label, total_mb, MAX_TOTAL_ATTACH_MB))

    delay_days = 0
    if suffix != "":
        try:
            delay_days = max(1, int(request.form.get("delay_days" + suffix) or 3))
        except ValueError:
            delay_days = 3

    # Signature + centrage des photos : propres à CHAQUE étape.
    want_sig = request.form.get("signature" + suffix) == "on"
    center = request.form.get("center_photos" + suffix) == "on"

    return {
        "label": label, "suffix": suffix,
        "subject": subject, "body": body, "delay_days": delay_days,
        "attachments": attachments, "inline_images": inline_images,
        "want_sig": want_sig, "center": center,
        "total_mb": round(total_mb, 2),
    }


def _photo_uris(inline_images):
    """Miniatures légères (data URI) pour l'aperçu."""
    out = {}
    for img in inline_images:
        uri = None
        try:
            from PIL import Image as _Image
            import io as _io2
            im = _Image.open(img["path"]); im.thumbnail((640, 640))
            buf = _io2.BytesIO(); im.convert("RGB").save(buf, "JPEG", quality=70)
            uri = "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
        except Exception:
            try:
                with open(img["path"], "rb") as fh:
                    b64 = base64.b64encode(fh.read()).decode()
                mime = mimetypes.guess_type(img["filename"])[0] or "image/png"
                uri = "data:%s;base64,%s" % (mime, b64)
            except OSError:
                continue
        out[img["cid"]] = uri
    return out


def _parse_start_at(raw):
    """Champ datetime-local ('YYYY-MM-DDTHH:MM') -> 'YYYY-MM-DD HH:MM:SS' (heure de
    Paris) si c'est bien dans le futur. Sinon '' (= démarrer dès que possible)."""
    raw = (raw or "").strip().replace("T", " ")
    if not raw:
        return ""
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            dt = datetime.strptime(raw, fmt)
            break
        except ValueError:
            dt = None
    if not dt:
        return ""
    now = warmup.now_paris().replace(tzinfo=None)
    if dt <= now:
        return ""                      # date passée -> on démarre tout de suite
    return dt.strftime("%Y-%m-%d %H:%M:%S")


@app.route("/preview", methods=["POST"])
@login_required
def preview():
    u = current_user()
    mode = request.form.get("mode", "campagne")
    if mode not in ("simple", "campagne"):
        mode = "campagne"
    campaign_name = (request.form.get("campaign_name") or "").strip()
    delay = max(0, int(request.form.get("delay") or DEFAULT_DELAY_S))
    start_at = _parse_start_at(request.form.get("start_at", ""))
    start_raw = (request.form.get("start_at") or "").strip()   # pour ré-afficher le champ
    recipients_raw = request.form.get("recipients", "")
    rows, errors = parse_recipients(recipients_raw)
    if not rows:
        errors.append("Aucun destinataire valide.")

    # --- Anti-doublon : exclure les prospects déjà contactés par certaines campagnes ---
    excl_sel = request.form.getlist("exclude_campaigns")   # ids de campagnes ou "ALL"
    excluded_count = 0
    if rows and excl_sel:
        if "ALL" in excl_sel:
            cids = warmup.all_campaign_ids()
        else:
            cids = [x for x in excl_sel if x and x != "ALL"]
        already = warmup.contacted_emails(cids)
        if already:
            before = len(rows)
            rows = [r for r in rows if (r.get("email", "").lower() not in already)]
            excluded_count = before - len(rows)
        if not rows:
            errors.append("Tous les destinataires ont déjà été contactés par les campagnes exclues — rien à envoyer.")

    user_sig = SIGNATURES.get(u["email"].lower(), "")
    tmpdir = tempfile.mkdtemp(prefix="relance_")

    # Étape 0 (obligatoire) + relances 1..5 (uniquement en mode campagne)
    steps = []
    step0 = _parse_step("", tmpdir, errors)
    if step0 is None:
        errors.append("Le premier mail (objet + corps) est obligatoire.")
    else:
        steps.append(step0)
    if mode == "campagne":
        for k in range(1, 6):
            st = _parse_step("_%d" % k, tmpdir, errors)
            if st is not None:
                steps.append(st)

    if not rows or not steps:
        for e in errors:
            flash(e)
        return redirect(url_for("index"))

    # Signature propre à CHAQUE étape (ex. off sur le 1er mail, on sur les relances).
    # Le centrage des photos est déjà lu par étape dans _parse_step.
    any_sig_missing = False
    for st in steps:
        st["signature_html"] = user_sig if st.get("want_sig") else ""
        if st.get("want_sig") and not user_sig:
            any_sig_missing = True
    if any_sig_missing:
        errors.append("Aucune signature configurée pour ton adresse — "
                      "les étapes où tu l'as cochée partiront sans signature.")

    # Garde-fou variables sur l'ensemble des étapes
    all_txt = " ".join((s["subject"] + " " + s["body"]) for s in steps)
    referenced = set(FIELD_RE.findall(all_txt))
    available = set().union(*[set(r.keys()) for r in rows]) if rows else set()
    unresolved = referenced - available
    if unresolved:
        errors.append("Variables non remplies (resteront telles quelles) : "
                      + ", ".join("{%s}" % v for v in sorted(unresolved))
                      + ". En liste d'adresses simple, seule {email} est connue.")

    # --- VÉRIF ANTI-COQUILLE : accolades déséquilibrées ---------------------
    # FIELD_RE ne matche que les variables VALIDES {xxx}. Une coquille comme
    # « objet_mail1} » (accolade ouvrante manquante) ou « {prenom » (fermante
    # manquante) échappe donc au contrôle ci-dessus : la « variable » part alors
    # EN CLAIR, non personnalisée (c'est le bug vécu sur l'objet du 1er mail).
    # On la détecte en retirant d'abord les {xxx} valides : s'il reste une
    # accolade, c'est qu'une variable est mal fermée.
    for _i, _s in enumerate(steps, 1):
        for _libelle, _val in (("l'objet", _s.get("subject", "")),
                               ("le corps", _s.get("body", ""))):
            _reste = FIELD_RE.sub("", _val or "")
            if "{" in _reste or "}" in _reste:
                errors.append(
                    "⚠️ Étape %d : accolade manquante ou en trop dans %s "
                    "(ex. « objet_mail1} » au lieu de « {objet_mail1} »). "
                    "Une variable mal fermée part EN CLAIR, sans être personnalisée — "
                    "corrige avant d'envoyer." % (_i, _libelle))

    # Aperçus : étape 0 sur les 5 premiers destinataires ; chaque relance sur le 1er
    PREVIEW_LIMIT = 5
    p0 = _photo_uris(step0["inline_images"])

    # Aperçu des images PAR PROSPECT : on télécharge à la volée (cache par URL) pour
    # les afficher, et on note les colonnes dont le lien est mort (souvent expiré).
    _img_cache = {}
    img_fail_cols = set()

    def _preview_img_resolver(col, url):
        if not re.match(r"^https?://", url or "", re.I):
            return ""
        if url not in _img_cache:
            dat = warmup.fetch_image_datauri(url, timeout=6)
            _img_cache[url] = dat or ""
            if not dat:
                img_fail_cols.add(col)
        return _img_cache[url]

    def render_for(step, row, photo_uris):
        html = assemble_body(step["body"], row, step["inline_images"],
                             step.get("signature_html", ""), step.get("center", False),
                             img_resolver=_preview_img_resolver)
        for cid, uri in photo_uris.items():
            html = html.replace("cid:%s" % cid, uri)
        return _links_new_tab(html)

    previews = [{"email": r["email"], "subject": personalize(step0["subject"], r),
                 "body": render_for(step0, r, p0)}
                for r in rows[:PREVIEW_LIMIT]]
    more_count = max(0, len(rows) - PREVIEW_LIMIT)

    # Résumé de la séquence (une carte par étape, rendue sur le 1er destinataire)
    r0 = rows[0]
    sequence = []
    for i, st in enumerate(steps):
        puris = p0 if i == 0 else _photo_uris(st["inline_images"])
        sequence.append({
            "index": i, "is_relance": i > 0,
            "label": st["label"], "delay_days": st["delay_days"],
            "subject": personalize(st["subject"], r0),
            "body": render_for(st, r0, puris),
            "n_photos": len(st["inline_images"]), "n_attach": len(st["attachments"]),
        })

    if img_fail_cols:
        errors.append("🖼️ Photo(s) non téléchargeable(s) — colonne(s) : "
                      + ", ".join("{%s}" % c for c in sorted(img_fail_cols))
                      + ". Le lien est probablement expiré (les liens Airtable expirent en quelques heures). "
                        "Ré-exporte depuis Airtable puis lance la campagne dans la foulée : "
                        "les photos valides seront téléchargées et intégrées au mail au lancement.")

    total_mb = round(sum(s["total_mb"] for s in steps), 2)

    job_id = uuid.uuid4().hex
    with LOCK:
        JOBS[job_id] = {
            "id": job_id, "status": "préparé",
            # étape 0 à plat (compat envoi immédiat "Envoyer la série")
            "subject": step0["subject"], "body": step0["body"],
            "attachments": step0["attachments"], "inline_images": step0["inline_images"],
            "signature_html": step0["signature_html"], "center": step0["center"],
            # la séquence complète (pour la chauffe)
            "steps": steps, "mode": mode, "name": campaign_name,
            "delay": delay, "start_at": start_at, "recipients": rows, "tmpdir": tmpdir,
            "recipients_raw": recipients_raw, "want_sig": step0["want_sig"],
            "sender_email": u["email"], "sender_password": None,
            "sender_name": u["name"],
            "results": [], "done": 0, "ok": 0, "ko": 0,
            "total": len(rows), "cancel": False,
        }
    est_min = round((len(rows) - 1) * delay / 60, 1) if len(rows) > 1 else 0
    heavy_pdf = next((a for a in step0["attachments"]
                      if a["filename"].lower().endswith(".pdf") and a["size_kb"] > 12 * 1024), None)
    return render_template("preview.html", job_id=job_id, previews=previews,
                           attachments=step0["attachments"], inline_images=step0["inline_images"],
                           total_mb=total_mb, delay=delay, est_min=est_min,
                           warnings=errors, count=len(rows), more_count=more_count,
                           sequence=sequence, n_steps=len(steps),
                           mode=mode, campaign_name=campaign_name,
                           excluded_count=excluded_count, start_at=start_at,
                           from_name=u["name"], from_email=u["email"], heavy_pdf=heavy_pdf)


@app.route("/send", methods=["POST"])
@login_required
def send():
    job_id = request.form.get("job_id", "")
    with LOCK:
        job = JOBS.get(job_id)
    if not job or job["status"] != "préparé":
        flash("Session d'envoi introuvable ou déjà lancée. Recommencez.")
        return redirect(url_for("index"))
    # Sécurité : on n'envoie que ses propres séries
    if job["sender_email"] != current_user()["email"]:
        flash("Cette série appartient à un autre expéditeur.")
        return redirect(url_for("index"))
    # Mot de passe (en mémoire) : perdu après un redémarrage -> reconnexion
    pwd = PASSWORDS.get(job["sender_email"])
    if not pwd:
        flash("Session de sécurité expirée (le service a redémarré). "
              "Reconnecte-toi puis relance l'envoi.")
        return redirect(url_for("login"))
    job["sender_password"] = pwd
    job["status"] = "running"
    threading.Thread(target=run_job, args=(job_id,), daemon=True).start()
    return redirect(url_for("progress_page", job_id=job_id))


def identity_required(f):
    """Comme login_required mais SANS exiger le mot de passe : sert à consulter
    la progression d'un envoi déjà lancé, même après un redémarrage."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not current_user():
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapper


@app.route("/progress/<job_id>")
@identity_required
def progress_page(job_id):
    with LOCK:
        job = JOBS.get(job_id)
    if not job:
        flash("Job introuvable."); return redirect(url_for("index"))
    return render_template("progress.html", job_id=job_id, total=job["total"])


@app.route("/progress/<job_id>/data")
@identity_required
def progress_data(job_id):
    with LOCK:
        job = JOBS.get(job_id)
    if not job:
        return jsonify({"status": "introuvable", "total": 0, "done": 0,
                        "ok": 0, "ko": 0, "results": []}), 200
    return jsonify({"status": job["status"], "total": job["total"], "done": job["done"],
                    "ok": job["ok"], "ko": job["ko"], "results": job["results"]})


@app.route("/cancel/<job_id>", methods=["POST"])
@identity_required
def cancel(job_id):
    with LOCK:
        job = JOBS.get(job_id)
    if job: job["cancel"] = True
    return jsonify({"ok": True})


@app.route("/journal")
@login_required
def journal():
    entries = []
    if os.path.exists(LOG_FILE):
        with open(LOG_FILE, encoding="utf-8") as fh:
            for line in fh:
                try: entries.append(json.loads(line))
                except json.JSONDecodeError: pass
    entries.reverse()
    return render_template("journal.html", entries=entries[:500], total=len(entries))


@app.route("/journal/download")
@login_required
def journal_download():
    if not os.path.exists(LOG_FILE):
        flash("Aucun journal pour le moment."); return redirect(url_for("journal"))
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from io import BytesIO

    entries = []
    with open(LOG_FILE, encoding="utf-8") as fh:
        for line in fh:
            try:
                entries.append(json.loads(line))
            except json.JSONDecodeError:
                pass
    entries.reverse()  # plus récent en haut

    wb = Workbook(); ws = wb.active; ws.title = "Journal des envois"
    headers = ["Date/heure", "Expéditeur", "Destinataire", "Objet",
               "Statut", "Copie Envoyés", "Mode", "Erreur"]
    ws.append(headers)
    hf = Font(name="Arial", bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="3C214B")
    thin = Side(style="thin", color="DDDDDD"); border = Border(thin, thin, thin, thin)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hf; cell.fill = hfill; cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    for e in entries:
        ws.append([
            e.get("ts", ""), e.get("expediteur", ""), e.get("email", ""),
            e.get("subject", ""), e.get("status", ""), e.get("copie", ""),
            e.get("mode", ""), e.get("error") or e.get("copie_err", ""),
        ])
    for i, w in enumerate([20, 30, 32, 42, 12, 14, 12, 42], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s1" % get_column_letter(len(headers))

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    fname = "journal_envois_%s.xlsx" % datetime.now().strftime("%Y-%m-%d")
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/healthz")
def healthz():
    """Signal vital. NE touche PAS la base (sinon il se figerait lui aussi) : il
    lit l'horodatage du dernier accès base réussi par le planificateur.

    Si la base devient injoignable (connexions qui se figent) ou si le thread du
    planificateur meurt, cet horodatage cesse d'avancer -> on répond 503 -> Render
    redémarre l'instance tout seul. Sans ça, le service reste « vivant » aux yeux
    de Render tout en étant totalement bloqué : c'est ce qui provoquait les gels
    sans fin (l'outil qui « tourne dans le vide » indéfiniment)."""
    # Filet de sécurité : si le planificateur n'est pas là (jamais démarré ou
    # disparu), on le relance immédiatement — sinon plus aucun mail ne partirait.
    try:
        warmup.ensure_scheduler()
    except Exception as e:
        print("[healthz] relance planificateur KO : %s" % e, flush=True)
    try:
        stale = warmup.db_stale_seconds()
    except Exception:
        return "ok", 200                      # sonde indisponible : on ne casse rien
    if stale > HEALTH_MAX_STALE_S:
        print("[healthz] base inaccessible depuis %ds -> redemarrage demande"
              % int(stale), flush=True)
        return "db stale (%ds)" % int(stale), 503
    return "ok", 200


@app.errorhandler(500)
def server_error(e):
    return (
        "<div style='font-family:Arial;max-width:560px;margin:60px auto;padding:0 20px;color:#2b2230;line-height:1.6'>"
        "<h1 style='color:#4A1E5C'>Oups, petit hoquet du serveur 🙈</h1>"
        "<p>Une opération n'a pas abouti — souvent parce que le serveur venait de redémarrer. "
        "<strong>Rien n'est perdu.</strong> Réessaie dans quelques secondes.</p>"
        "<p style='margin-top:18px'><a href='/' style='color:#B90009;font-weight:bold;text-decoration:none'>← Retour à l'accueil</a></p>"
        "</div>", 500)


@app.route("/diag/threads")
@login_required
def diag_threads():
    """Diagnostic : pile d'appel de CHAQUE thread. Permet d'identifier un blocage
    (quel thread attend quoi, et à quelle ligne). Lecture seule, sans effet."""
    import sys as _sys
    import traceback as _tb
    import threading as _th
    noms = {t.ident: t.name for t in _th.enumerate()}
    out = []
    for tid, frame in _sys._current_frames().items():
        out.append("=== thread %s (%s) ===\n" % (tid, noms.get(tid, "?")))
        out.extend(_tb.format_stack(frame))
        out.append("\n")
    return "<pre>" + "".join(out) + "</pre>"


@app.errorhandler(413)
def too_large(e):
    flash(f"Envoi trop volumineux : le formulaire dépasse "
          f"{int(MAX_TOTAL_ATTACH_MB) + 20} Mo au total (pièces jointes + liste de "
          f"prospects collée). Allège les pièces jointes, ou découpe ta liste en deux "
          f"campagnes (rappel : au-delà de ~25 Mo, beaucoup de boîtes rejettent le mail).")
    return redirect(url_for("index"))


# --------------------------------------------------------------------------- #
# CHAUFFE (warm-up) : envoi unitaire + démarrage du planificateur
# --------------------------------------------------------------------------- #
def _warmup_send_one(sender_email, password, camp, row):
    """Construit + envoie UN mail pour la chauffe, range une copie dans Envoyés,
    et journalise. Réutilise tout le code MIME existant ci-dessus.
    Renvoie (ok: bool, message_id: str|None, err: str|None)."""
    from email.utils import make_msgid
    to_email = row.get("email", "")
    try:
        subject = personalize(camp["subject"], row)
        # Image par prospect téléchargée au lancement -> intégrée en inline (CID),
        # pour qu'elle s'affiche partout même si le lien d'origine a expiré depuis.
        imgfiles = row.get("__imgfiles") or {}
        extra_inline = []
        def _img_resolver(col, url, _files=imgfiles, _acc=extra_inline):
            p = _files.get(col)
            if p and os.path.exists(p):
                cid = "rimg%s" % uuid.uuid4().hex[:12]
                _acc.append({"path": p, "cid": cid, "filename": os.path.basename(p)})
                return "cid:" + cid
            return url if re.match(r"^https?://", url or "", re.I) else ""
        body = assemble_body(camp["body"], row, camp.get("inline_images", []),
                             camp.get("signature_html", ""), camp.get("center", False),
                             img_resolver=_img_resolver)
        inline_all = list(camp.get("inline_images", [])) + extra_inline
        msg = build_message(sender_email, camp.get("sender_name", ""), to_email, subject,
                            body, camp.get("attachments", []), inline_all)
        # Message-ID propre (sur le domaine de l'expéditeur) -> matching des réponses
        msgid = make_msgid(domain=sender_email.split("@")[-1])
        msg["Message-ID"] = msgid
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as server:
            server.ehlo(); server.starttls(); server.ehlo()
            server.login(sender_email, password)
            server.send_message(msg)
        # Copie "Envoyés" (best-effort, n'échoue jamais l'envoi)
        copied, copie_err = False, None
        try:
            copied, copie_err = save_to_sent(sender_email, password, msg)
        except Exception as e:
            copied, copie_err = False, str(e)
        if not copied:
            print("[copie Envoyés KO] %s -> %s : %s" % (sender_email, to_email, copie_err),
                  flush=True)
        log_send({"email": to_email, "subject": subject, "expediteur": sender_email,
                  "status": "OK", "copie": "OK" if copied else "KO",
                  "copie_err": copie_err or "", "mode": "chauffe"})
        return True, msgid, None
    except Exception as e:
        log_send({"email": to_email, "subject": camp.get("subject", ""),
                  "expediteur": sender_email, "status": "ERREUR",
                  "error": str(e), "mode": "chauffe"})
        return False, None, str(e)


warmup.init_app(
    app,
    get_password=_password_for,
    send_one=_warmup_send_one,
    get_job=lambda jid: JOBS.get(jid),
    imap_cfg={"host": IMAP_HOST, "port": IMAP_PORT},
)


# --------------------------------------------------------------------------- #
# Brique « Levées +12 mois » — déclenchement MANUEL (test / rattrapage)
# --------------------------------------------------------------------------- #
# Le run hebdo automatique NE dépend PAS de cette route : il est accroché dans
# warmup._loop() (moteur de chauffe), donc un rebuild de ce app.py ne peut plus
# désactiver l'automatisation. Cette route ne sert qu'aux runs à la demande :
#   .../chauffe/levees/run-now             → fenêtre normale (LEVEES_WINDOW_DAYS)
#   .../chauffe/levees/run-now?backfill=1  → tout le stock +12 mois (rattrapage)
@app.route("/chauffe/levees/run-now", methods=["GET", "POST"])
@login_required
def levees_run_now():
    backfill = request.args.get("backfill") in ("1", "true", "yes")
    if levees_auto.run_now_async(backfill=backfill):
        flash("✅ Run Levées lancé%s. Suis l'avancement dans les logs Render / le journal."
              % (" — BACKFILL (tout le stock +12 mois)" if backfill else ""))
    else:
        flash("Un run Levées est déjà en cours — réessaie dans quelques minutes.")
    return redirect(url_for("warmup.dashboard"))


# NE PAS démarrer le planificateur à l'import : gunicorn importe le module dans le
# process maître ET dans le worker, ce qui lançait DEUX moteurs (d'où les envois
# partant par paires). On le démarre paresseusement, uniquement dans le process qui
# sert vraiment des requêtes (le worker) — le maître, lui, n'en lancera jamais.
# ensure_scheduler() est idempotent (verrou + vérif « déjà vivant »), et /healthz
# l'appelle déjà toutes les ~5 s : le moteur démarre donc dans les secondes qui
# suivent le boot, et UN SEUL tourne.
@app.before_request
def _demarrer_planificateur_si_besoin():
    warmup.ensure_scheduler()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)
